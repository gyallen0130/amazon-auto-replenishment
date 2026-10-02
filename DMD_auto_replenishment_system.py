#!/usr/bin/env python3
# DMD automatic replenishment decision system v1.4.3 (stale inventory + Business Report + cart warning + visual alerts)
import os, re, math, time, getpass
from datetime import datetime, timedelta, timezone
from pathlib import Path
import requests
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.utils import get_column_letter

DMD_PREFIX = 'DMD'
KEEPA_DOMAIN = 5
OFFERS_COUNT = 20
MIN_TOKENS_LEFT = 35
FRESH_MINUTES = 180
DEFAULT_LEAD_DAYS = 7
DEFAULT_SAFETY_DAYS = 7
MIN_PROFIT_MARGIN = 0.0  # 売上拡大モード: 利益率0%以上を仕入れ候補
OFFERS_COVER_DAYS = 30
BUSINESS_MIN_SESSIONS = 30
CART_RED_PCT = 10.0
CART_ORANGE_PCT = 30.0
CART_YELLOW_PCT = 50.0
MASTER_PATH = '/content/drive/MyDrive/02_AI業務効率化/AI仕入れ判断/product_master.xlsm'
RESULT_DIR = '/content/drive/MyDrive/02_AI業務効率化/AI仕入れ判断/01_出力結果'


def _looks_like_header(row):
    """Amazonの説明行を飛ばし、本物の表ヘッダーらしい行を検出する。"""
    vals = {str(x).strip() for x in row if x is not None and str(x).strip()}
    sales_score = sum(x in vals for x in ('日付/時間','トランザクションの種類','SKU','数量'))
    inventory_score = sum(x in vals for x in ('出品者SKU','ASIN','販売できる数量'))
    master_score = sum(x in vals for x in ('SKU','ASIN','商品名','仕入単価(税抜)'))
    business_score = sum(x in vals for x in ('SKU','（子）ASIN','セッション数 - 合計','おすすめ出品（おすすめ商品）の獲得率'))
    return sales_score >= 3 or inventory_score >= 2 or master_score >= 3 or business_score >= 3


def normalize_table(rows):
    """先頭のタイトル・説明文・空行を除去し、ヘッダー行から始まる表に正規化。"""
    if not rows:
        return rows
    for i, row in enumerate(rows[:100]):
        if _looks_like_header(row):
            return rows[i:]
    return rows


def read_tabular(path):
    path = str(path)
    suffix = Path(path).suffix.lower()
    if suffix == '.csv':
        import csv
        last_error = None
        for enc in ('utf-8-sig','cp932','utf-8'):
            try:
                with open(path, encoding=enc, newline='') as f:
                    rows = list(csv.reader(f))
                return normalize_table(rows)
            except UnicodeDecodeError as e:
                last_error = e
                continue
        raise ValueError(f'CSVの文字コードを読めませんでした: {path}') from last_error

    if suffix in ('.xlsx','.xlsm'):
        wb = load_workbook(path, data_only=True, read_only=True, keep_vba=(suffix == '.xlsm'))
        # 先頭シートだけに固定せず、認識可能な表を持つ最初のシートを使用。
        fallback = None
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            if fallback is None:
                fallback = rows
            normalized = normalize_table(rows)
            if normalized and _looks_like_header(normalized[0]):
                return normalized
        return normalize_table(fallback or [])

    raise ValueError(f'未対応のファイル形式です: {suffix}')


def rows_as_dicts(rows):
    rows = normalize_table(rows)
    if not rows:
        return []
    headers=[str(x).strip() if x is not None else '' for x in rows[0]]
    return [dict(zip(headers,r)) for r in rows[1:] if any(v is not None and str(v).strip() != '' for v in r)]


def parse_dt(v):
    if isinstance(v, datetime): return v.replace(tzinfo=None)
    s=str(v or '').replace(' JST','').strip()
    for fmt in ('%Y/%m/%d %H:%M:%S','%Y-%m-%d %H:%M:%S','%Y/%m/%d','%Y-%m-%d'):
        try: return datetime.strptime(s,fmt)
        except: pass
    return None


def num(v, default=0.0):
    try:
        if v is None or v == '' or str(v).upper() in ('#N/A','N/A','NONE'): return default
        s=str(v).strip().replace(',', '').replace('￥','').replace('¥','')
        if s.endswith('%'): return float(s[:-1]) / 100.0
        return float(s)
    except: return default

def pct_value(v):
    """Business Reportの%文字列を0〜100の百分率へ。例: '27.00%' -> 27.0"""
    if v is None or str(v).strip()=='' : return None
    s=str(v).strip().replace(',', '')
    try:
        return float(s[:-1]) if s.endswith('%') else float(s) * (100.0 if float(s) <= 1 else 1.0)
    except:
        return None


def ceil_unit(qty, minimum=1, unit=1, case=1):
    if qty <= 0: return 0
    step=max(1,int(num(unit,1)),int(num(case,1)))
    q=max(int(math.ceil(qty)),int(num(minimum,1)))
    return int(math.ceil(q/step)*step)


def keepa_minutes_to_dt(m):
    if m is None or m < 0: return None
    return datetime(2011,1,1) + timedelta(minutes=int(m))


def keepa_get(api_key, asin, offers=False):
    params={'key':api_key,'domain':KEEPA_DOMAIN,'asin':asin,'stats':90,'history':1}
    if offers: params['offers']=OFFERS_COUNT
    r=requests.get('https://api.keepa.com/product',params=params,timeout=60)
    r.raise_for_status()
    return r.json()


def wait_for_tokens(api_key, needed=MIN_TOKENS_LEFT):
    try:
        r=requests.get('https://api.keepa.com/token',params={'key':api_key},timeout=30).json()
        left=int(r.get('tokensLeft',0)); refill=int(r.get('refillRate',5) or 5)
        if left < needed:
            secs=math.ceil((needed-left)/max(refill,1))*60+3
            print(f'Keepaトークン待機: {left} -> {needed} / 約{secs//60}分')
            time.sleep(secs)
    except Exception as e:
        print('トークン確認をスキップ:',e)


def keepa_summary(product):
    stats=product.get('stats') or {}
    current=stats.get('current') or []
    avg30=stats.get('avg30') or []
    avg90=stats.get('avg90') or []
    def price(arr, idx):
        try:
            v=arr[idx]
            # Keepa domain=5 (Amazon.co.jp): price integer is handled as JPY.
            # v1.2 divided by 100, which made the output sales price two digits too small.
            return None if v is None or v < 0 else float(v)
        except:
            return None
    def raw(arr, idx):
        try:
            v=arr[idx]
            return None if v is None or v < 0 else v
        except: return None
    # Keepa CSV indices: AMAZON=0, NEW=1, SALES=3, BUY_BOX_SHIPPING=18 (commonly used)
    return {
        'keepa_title': product.get('title'),
        'amazon_price': price(current,0),
        'new_price': price(current,1),
        'new_avg30': price(avg30,1),
        'new_avg90': price(avg90,1),
        'rank_now': raw(current,3),
        'rank_avg30': raw(avg30,3),
        'rank_avg90': raw(avg90,3),
        'amazon_present': 'あり' if price(current,0) else 'なし',
    }


def offer_summary(product):
    offers=product.get('offers') or []
    live=set(product.get('liveOffersOrder') or [])
    now=datetime.utcnow()
    active=[]
    for i,o in enumerate(offers):
        if o.get('condition') not in (0,None): continue
        if o.get('isShippable') is False: continue
        seen=keepa_minutes_to_dt(o.get('lastSeen'))
        fresh=(seen is not None and (now-seen).total_seconds() <= FRESH_MINUTES*60)
        if i not in live and not fresh: continue
        csv=o.get('offerCSV') or []
        p=None
        if len(csv)>=2:
            p=num(csv[-2],-1); ship=num(csv[-1],0)
            if p>=0:
                # Amazon.co.jp (domain=5): offer price + shipping are JPY values.
                p=float(p+max(ship,0))
            else:
                p=None
        active.append((o,p))
    prices=[p for _,p in active if p is not None]
    fba=[p for o,p in active if o.get('isFBA') and p is not None]
    fbm=[p for o,p in active if not o.get('isFBA') and p is not None]
    amz=any(o.get('isAmazon') for o,_ in active)
    return {'offer_count':len(active),'fba_count':sum(1 for o,_ in active if o.get('isFBA')),
            'fbm_count':sum(1 for o,_ in active if not o.get('isFBA')),
            'amazon_offers':'あり' if amz else 'なし','lowest_new':min(prices) if prices else None,
            'lowest_fba':min(fba) if fba else None,'lowest_fbm':min(fbm) if fbm else None}


def classify_input_file(path):
    """列名を見て sales / inventory / master を判定。xlsx/xlsm/csv 対応。"""
    try:
        rows = read_tabular(path)
        if not rows:
            return None
        headers = {str(x).strip() for x in rows[0] if x is not None}
    except Exception as e:
        print(f'ファイル確認失敗: {path} / {e}')
        return None

    # 商品マスタを最優先。ファイル名が変わっても列名で判定する。
    if {'SKU','ASIN','商品名','仕入単価(税抜)'}.issubset(headers):
        return 'master'
    # Amazon FBA在庫レポート
    if ('出品者SKU' in headers or 'SKU' in headers) and ('販売できる数量' in headers or '現在庫' in headers):
        return 'inventory'
    # Amazonトランザクション/売上レポート
    if {'トランザクションの種類','SKU','数量'}.issubset(headers) and ('日付/時間' in headers or '日付' in headers):
        return 'sales'
    # ビジネスレポート：詳細ページ 売上・トラフィック（子商品別）
    if {'SKU','（子）ASIN','セッション数 - 合計','おすすめ出品（おすすめ商品）の獲得率'}.issubset(headers):
        return 'business'
    return None


def check_google_drive():
    """v1.4.2.1: Driveの認証UIは.py内から呼ばず、Colabセル側で事前マウントする。"""
    drive_root = Path('/content/drive')
    if not drive_root.exists():
        raise RuntimeError(
            "Google Driveがマウントされていません。Colabの別セルで\n"
            "from google.colab import drive\n"
            "drive.mount('/content/drive')\n"
            "を実行してから、このプログラムを実行してください。"
        )
    print('Google Drive接続: OK（マウント済み）')


def choose_files_interactive():
    """v1.4.2: /content直下の売上・在庫ファイルを自動検出する。"""
    base = Path('/content')
    supported = ('.xlsx', '.xlsm', '.csv')
    names = [str(p) for p in base.iterdir()
             if p.is_file() and p.suffix.lower() in supported
             and p.name != 'product_master.xlsm'
             and not p.name.startswith('DMD仕入れ判断_')]

    # 新しいファイルを優先
    names.sort(key=lambda n: Path(n).stat().st_mtime, reverse=True)

    found = {'inventory': [], 'sales': [], 'business': []}
    for n in names:
        try:
            kind = classify_input_file(n)
        except Exception:
            continue
        if kind in found:
            found[kind].append(n)

    # 列判定で拾えなかった場合のみファイル名を補助的に使用
    if not found['inventory']:
        found['inventory'] = [n for n in names
                              if '在庫' in Path(n).name or 'inventory' in Path(n).name.lower()
                              or 'zaiko' in Path(n).name.lower()]
    if not found['sales']:
        found['sales'] = [n for n in names
                          if 'transaction' in Path(n).name.lower() or '売上' in Path(n).name
                          or 'sales' in Path(n).name.lower() or 'uriage' in Path(n).name.lower()]
    # Business Reportは似た名前の日別レポートがあるため、ファイル名だけでは補完しない。
    # 子商品別の必須列が揃うファイルだけを採用する。

    result = {
        'inventory': found['inventory'][0] if found['inventory'] else None,
        'sales': found['sales'][0] if found['sales'] else None,
        'business': found['business'][0] if found['business'] else None,
    }
    print('自動判別結果:', {'master': MASTER_PATH, **result})

    missing = []
    if result['sales'] is None: missing.append('売上CSV')
    if result['inventory'] is None: missing.append('在庫CSV')
    if missing:
        raise FileNotFoundError(
            '・'.join(missing) + ' を /content 直下から検出できませんでした。\n'
            'Colab左側のファイル欄へ売上CSV・在庫CSVをアップロードしてから再実行してください。'
        )

    print('売上CSV:', Path(result['sales']).name)
    print('在庫CSV:', Path(result['inventory']).name)
    if result['business']:
        print('ビジネスレポート:', Path(result['business']).name)
    else:
        print('ビジネスレポート: 未検出（カート関連列は空欄で続行）')
    return result['sales'], result['inventory'], result['business']


def unique_output_path(result_dir, base_name):
    """既存ファイルを上書きせず _02, _03... と採番。"""
    os.makedirs(result_dir, exist_ok=True)
    p = Path(result_dir) / base_name
    if not p.exists(): return str(p)
    stem, suffix = p.stem, p.suffix
    i = 2
    while True:
        candidate = p.with_name(f'{stem}_{i:02d}{suffix}')
        if not candidate.exists(): return str(candidate)
        i += 1

def main(sales_path=None, inventory_path=None, business_path=None, master_path=MASTER_PATH, api_key=None, output_path=None):
    check_google_drive()
    if not os.path.exists(master_path):
        raise FileNotFoundError(f'商品マスタが見つかりません: {master_path}')
    print(f'商品マスタ: OK\n  {master_path}')
    if not all((sales_path, inventory_path)):
        sales_path,inventory_path,business_path=choose_files_interactive()
    if api_key is None:
        api_key=getpass.getpass('Keepa APIキーを入力（画面には表示されません）: ').strip()

    sales=rows_as_dicts(read_tabular(sales_path))
    inv=rows_as_dicts(read_tabular(inventory_path))
    master=rows_as_dicts(read_tabular(master_path))
    business=rows_as_dicts(read_tabular(business_path)) if business_path else []
    master=[r for r in master if str(r.get('SKU') or '').startswith(DMD_PREFIX) and str(r.get('有効') or '有効')!='停止']

    dates=[parse_dt(r.get('日付/時間')) for r in sales if str(r.get('トランザクションの種類') or '')=='注文']
    dates=[d for d in dates if d]
    if not dates: raise ValueError('売上ファイルから注文日を取得できません。')
    end=max(dates); start30=end-timedelta(days=29); start90=end-timedelta(days=89)

    s30={}; s90={}; last_sale={}
    for r in sales:
        sku=str(r.get('SKU') or '')
        if not sku.startswith(DMD_PREFIX) or str(r.get('トランザクションの種類') or '')!='注文': continue
        d=parse_dt(r.get('日付/時間')); q=num(r.get('数量'),0)
        if not d: continue
        if q > 0 and (sku not in last_sale or d > last_sale[sku]): last_sale[sku]=d
        if d>=start90: s90[sku]=s90.get(sku,0)+q
        if d>=start30: s30[sku]=s30.get(sku,0)+q

    stock={}; asin_inv={}
    for r in inv:
        sku=str(r.get('出品者SKU') or r.get('SKU') or '')
        if not sku.startswith(DMD_PREFIX): continue
        stock[sku]=stock.get(sku,0)+num(r.get('販売できる数量') if '販売できる数量' in r else r.get('現在庫'),0)
        if r.get('ASIN'): asin_inv[sku]=r.get('ASIN')

    # Business ReportはSKUで直接結合。DMD以外は対象外。
    business_by_sku={}
    for r in business:
        sku=str(r.get('SKU') or '').strip()
        if not sku.startswith(DMD_PREFIX): continue
        business_by_sku[sku]={
            'セッション数': int(num(r.get('セッション数 - 合計'),0)),
            'カート獲得率': pct_value(r.get('おすすめ出品（おすすめ商品）の獲得率')),
            'BR注文点数': int(num(r.get('注文された商品点数'),0)),
            'BRユニットセッション率': pct_value(r.get('ユニットセッション率')),
        }

    records=[]
    print(f'DMD商品マスタ: {len(master)} SKU / 売上基準日: {end.date()}')
    # First pass: basic Keepa data for all master SKUs.
    for idx,m in enumerate(master,1):
        sku=str(m.get('SKU')); asin=str(m.get('ASIN') or asin_inv.get(sku) or '')
        rec={'SKU':sku,'ASIN':asin,'商品名':m.get('商品名'),'現在庫':stock.get(sku,0),
             '30日販売':s30.get(sku,0),'90日販売':s90.get(sku,0),'master':m}
        rec.update(business_by_sku.get(sku, {'セッション数':None,'カート獲得率':None,'BR注文点数':None,'BRユニットセッション率':None}))
        ls=last_sale.get(sku)
        rec['最終販売日']=ls.date().isoformat() if ls else '90日以内なし'
        rec['無販売日数']=(end.date()-ls.date()).days if ls else 90
        if rec['現在庫'] > 0:
            if ls is None or rec['無販売日数'] >= 90: rec['滞留アラート']='90日以上販売なし'
            elif rec['無販売日数'] >= 60: rec['滞留アラート']='60日以上販売なし'
            elif rec['無販売日数'] >= 30: rec['滞留アラート']='30日以上販売なし'
            else: rec['滞留アラート']=''
        else:
            rec['滞留アラート']=''
        if asin:
            try:
                wait_for_tokens(api_key,3)
                js=keepa_get(api_key,asin,False); p=(js.get('products') or [{}])[0]
                rec.update(keepa_summary(p)); time.sleep(.25)
            except Exception as e: rec['keepa_error']=str(e)
        records.append(rec)
        if idx%20==0: print(f'Keepa基本取得 {idx}/{len(master)}')

    # Demand forecast first, then Offers only for near-replenishment items.
    for r in records:
        s30v=r['30日販売']; s90v=r['90日販売']
        if s90v>0:
            base=.70*(s30v*1.5)+.30*(s90v*.5)
            r['45日需要予測']=max(0,round(base)); r['予測信頼度']='中' if s90v>=20 else '低'
            r['需要方式']='自社販売実績'
        else:
            r['45日需要予測']=0; r['予測信頼度']='参考'; r['需要方式']='販売実績なし（自動発注しない）'
        daily=r['45日需要予測']/45 if r['45日需要予測'] else 0
        r['在庫カバー日数']=r['現在庫']/daily if daily>0 else None

        # カート獲得率は単独で警告しない。直近30日の販売失速と組み合わせる。
        sessions=r.get('セッション数'); cart=r.get('カート獲得率')
        prior60=max(0, r['90日販売']-r['30日販売'])
        prior30pace=prior60/2.0
        r['前60日月換算販売']=round(prior30pace,1)
        if prior30pace > 0:
            r['販売速度比']=r['30日販売']/prior30pace
        else:
            r['販売速度比']=None
        alert=''
        if sessions is None or cart is None:
            alert='データなし'
        elif sessions < BUSINESS_MIN_SESSIONS:
            alert='データ不足'
        elif prior30pace > 0:
            ratio=r['販売速度比']
            if cart < CART_RED_PCT and ratio < 0.70:
                alert='異常：低カート＋販売大幅低下'
            elif cart < CART_ORANGE_PCT and ratio < 0.80:
                alert='要注意：低カート＋販売低下'
            elif cart < CART_YELLOW_PCT and ratio < 0.90:
                alert='注意：カート低下＋販売やや低下'
            else:
                alert='問題なし'
        else:
            # 比較期間に販売がない場合、カート0%でも現在売れていれば警告しない。
            alert='問題なし' if r['30日販売'] > 0 else ('要確認：販売実績なし' if cart < CART_ORANGE_PCT else '問題なし')
        r['カート獲得アラート']=alert

    offer_targets=[r for r in records if r['90日販売']>0 and (r['現在庫']==0 or (r['在庫カバー日数'] is not None and r['在庫カバー日数']<=OFFERS_COVER_DAYS))]
    print(f'Offers取得対象: {len(offer_targets)} SKU（在庫カバー{OFFERS_COVER_DAYS}日以下）')
    for idx,r in enumerate(offer_targets,1):
        if not r['ASIN']: continue
        try:
            wait_for_tokens(api_key,MIN_TOKENS_LEFT)
            js=keepa_get(api_key,r['ASIN'],True); p=(js.get('products') or [{}])[0]
            r.update(offer_summary(p)); time.sleep(.5)
        except Exception as e: r['offers_error']=str(e)
        if idx%10==0: print(f'Offers取得 {idx}/{len(offer_targets)}')

    # Final economics and replenishment decision.
    for r in records:
        m=r['master']; forecast=r['45日需要予測']; daily=forecast/45 if forecast else 0
        lead=num(m.get('標準納期(日)'),DEFAULT_LEAD_DAYS); safety=num(m.get('安全在庫日数'),DEFAULT_SAFETY_DAYS)
        target_days=lead+safety
        target_stock=math.ceil(daily*target_days) if daily else 0
        shortage=max(0,target_stock-r['現在庫'])
        r['目標在庫日数']=target_days; r['目標在庫数']=target_stock
        r['推奨発注数']=ceil_unit(shortage,m.get('最小発注数'),m.get('発注単位'),m.get('ケース入数')) if r['90日販売']>0 else 0

        # v1.4.2: 採算計算は現在のFBA最安値を最優先。FBA不在時のみ新品最安値へフォールバック。
        if r.get('lowest_fba') is not None:
            sale_price=r.get('lowest_fba'); price_type='FBA最安値'
        elif r.get('lowest_new') is not None:
            sale_price=r.get('lowest_new'); price_type='新品最安値（FBA不在）'
        elif r.get('new_price') is not None:
            sale_price=r.get('new_price'); price_type='Keepa新品価格（Offers未取得）'
        else:
            sale_price=None; price_type='価格取得不能'
        if sale_price is not None:
            sale_price=num(sale_price,0)
            if sale_price <= 0:
                sale_price=None; price_type='価格取得不能'
        cost=num(m.get('仕入単価(税抜)'),0)*(1+num(m.get('消費税率'),0))
        fee_rate=num(m.get('販売手数料率'),0); fba=num(m.get('FBA送料'),-1)
        r['販売価格']=sale_price; r['価格種別']=price_type; r['税込仕入原価']=cost
        r['現在庫金額']=round(cost * r['現在庫'])
        if sale_price and fba>=0:
            profit=sale_price-cost-sale_price*fee_rate-fba
            roi=profit/cost if cost>0 else None
            r['1個利益']=round(profit); r['ROI']=roi; r['利益率']=profit/sale_price if sale_price else None
            profitable=((profit/sale_price) >= MIN_PROFIT_MARGIN) if sale_price else False
            margin=r['利益率']
            r['利益区分']='積極補充' if margin>=0.10 else '通常補充' if margin>=0.05 else '売上重視補充' if margin>=0.02 else '売上拡大型' if margin>=0 else '赤字'
            r['45日予測売上高']=round(sale_price * forecast)
        else:
            r['1個利益']=None; r['ROI']=None; r['利益率']=None; r['利益区分']='要確認'; r['45日予測売上高']=None; profitable=None

        cov=r['在庫カバー日数']
        if r['90日販売']<=0:
            priority='D 再テスト候補' if (r.get('rank_now') or 10**9)<50000 else 'D 見送り'
            decision='自動発注しない'
        elif r['現在庫']==0: priority='1 最優先（欠品中）'; decision='発注' if profitable else ('利益要確認' if profitable is None else '見送り（利益不足）')
        elif cov is not None and cov<=7: priority='1 最優先（7日以内）'; decision='発注' if profitable else ('利益要確認' if profitable is None else '見送り（利益不足）')
        elif cov is not None and cov<=14: priority='2 至急（14日以内）'; decision='発注' if profitable else ('利益要確認' if profitable is None else '見送り（利益不足）')
        elif cov is not None and cov<=21: priority='3 早め（21日以内）'; decision='発注' if profitable else ('利益要確認' if profitable is None else '見送り（利益不足）')
        elif cov is not None and cov<=30: priority='4 注意（30日以内）'; decision='発注準備' if profitable else ('利益要確認' if profitable is None else '見送り（利益不足）')
        else: priority='5 当面不要'; decision='補充不要'
        # 在庫が残ったまま30日以上売れていないSKUは追加補充しない。
        if r.get('滞留アラート'):
            priority='5 当面不要（滞留）'
            decision='補充不要（滞留在庫）'
        r['補充優先度']=priority; r['仕入判断']=decision
        if decision.startswith('見送り') or decision=='利益要確認' or decision.startswith('補充不要'): r['推奨発注数']=0

    # Output workbook.
    wb=Workbook(); ws=wb.active; ws.title='仕入れ指示'
    headers=['補充優先度','仕入判断','SKU','ASIN','商品名','現在庫','現在庫金額','30日販売','90日販売','最終販売日','無販売日数','滞留アラート',
             'セッション数','カート獲得率','販売速度比','カート獲得アラート','45日需要予測','在庫カバー日数',
             '目標在庫日数','目標在庫数','推奨発注数','販売価格','価格種別','45日予測売上高','税込仕入原価','1個利益','利益率','ROI','利益区分','FBA競合数','Amazon本体','予測信頼度','需要方式']
    order={'1':1,'2':2,'3':3,'4':4,'5':5,'D':6}
    records.sort(key=lambda r:(order.get(str(r['補充優先度'])[0],9), r['在庫カバー日数'] if r['在庫カバー日数'] is not None else 99999,-r['45日需要予測']))
    ws.append(headers)
    for r in records:
        ws.append([r.get('補充優先度'),r.get('仕入判断'),r.get('SKU'),r.get('ASIN'),r.get('商品名'),r.get('現在庫'),r.get('現在庫金額'),r.get('30日販売'),r.get('90日販売'),r.get('最終販売日'),('90日以上' if r.get('最終販売日')=='90日以内なし' else r.get('無販売日数')),r.get('滞留アラート'),r.get('セッション数'),(r.get('カート獲得率')/100 if r.get('カート獲得率') is not None else None),(r.get('販売速度比')),r.get('カート獲得アラート'),r.get('45日需要予測'),r.get('在庫カバー日数'),r.get('目標在庫日数'),r.get('目標在庫数'),r.get('推奨発注数'),r.get('販売価格'),r.get('価格種別'),r.get('45日予測売上高'),r.get('税込仕入原価'),r.get('1個利益'),r.get('利益率'),r.get('ROI'),r.get('利益区分'),r.get('fba_count'),r.get('amazon_offers') or r.get('amazon_present'),r.get('予測信頼度'),r.get('需要方式')])
    blue='1F4E78'; thin=Side(style='thin',color='D9E1F2')
    for c in ws[1]: c.fill=PatternFill('solid',fgColor=blue); c.font=Font(bold=True,color='FFFFFF'); c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    for row in ws.iter_rows():
        for c in row: c.border=Border(left=thin,right=thin,top=thin,bottom=thin); c.alignment=Alignment(vertical='top',wrap_text=True)
    widths=[22,20,25,14,55,10,14,10,10,14,12,20,12,12,12,28,12,14,12,12,12,12,22,16,14,12,10,10,16,10,11,11,24]
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes='F2'; ws.auto_filter.ref=ws.dimensions
    col={name:i+1 for i,name in enumerate(headers)}
    colors={'darkgreen':'548235','green':'70AD47','lightgreen':'C6E0B4','yellow':'FFD966','orange':'F4B183','red':'C00000','gray':'D9E1F2'}
    def paint(cell, color, white=False):
        cell.fill=PatternFill('solid',fgColor=colors[color]); cell.font=Font(bold=True,color=('FFFFFF' if white else '000000'))
    for row in range(2,ws.max_row+1):
        ws.cell(row,col['在庫カバー日数']).number_format='0.0"日"'
        for name in ('現在庫金額','販売価格','45日予測売上高','税込仕入原価','1個利益'):
            ws.cell(row,col[name]).number_format='#,##0"円"'
        for name in ('カート獲得率','販売速度比','利益率','ROI'):
            ws.cell(row,col[name]).number_format='0.0%'

        # 補充優先度：仕入れる方向を緑、注意を黄、不要を灰色。
        p=str(ws.cell(row,col['補充優先度']).value or '')
        if p.startswith('1'): paint(ws.cell(row,col['補充優先度']),'darkgreen',True)
        elif p.startswith('2'): paint(ws.cell(row,col['補充優先度']),'green')
        elif p.startswith('3'): paint(ws.cell(row,col['補充優先度']),'lightgreen')
        elif p.startswith('4'): paint(ws.cell(row,col['補充優先度']),'yellow')
        else: paint(ws.cell(row,col['補充優先度']),'gray')

        d=str(ws.cell(row,col['仕入判断']).value or '')
        if d=='発注': paint(ws.cell(row,col['仕入判断']),'darkgreen',True)
        elif d=='発注準備': paint(ws.cell(row,col['仕入判断']),'lightgreen')
        elif '利益要確認' in d: paint(ws.cell(row,col['仕入判断']),'orange')
        elif d.startswith('見送り'): paint(ws.cell(row,col['仕入判断']),'red',True)
        else: paint(ws.cell(row,col['仕入判断']),'gray')

        profit_class=str(ws.cell(row,col['利益区分']).value or '')
        if profit_class=='積極補充': paint(ws.cell(row,col['利益区分']),'darkgreen',True)
        elif profit_class=='通常補充': paint(ws.cell(row,col['利益区分']),'lightgreen')
        elif profit_class=='売上重視補充': paint(ws.cell(row,col['利益区分']),'yellow')
        elif profit_class=='売上拡大型': paint(ws.cell(row,col['利益区分']),'orange')
        elif profit_class=='赤字': paint(ws.cell(row,col['利益区分']),'red',True)
        else: paint(ws.cell(row,col['利益区分']),'gray')

        stale=str(ws.cell(row,col['滞留アラート']).value or '')
        if stale.startswith('90'): paint(ws.cell(row,col['滞留アラート']),'red',True)
        elif stale.startswith('60'): paint(ws.cell(row,col['滞留アラート']),'orange')
        elif stale.startswith('30'): paint(ws.cell(row,col['滞留アラート']),'yellow')

        cart_alert=str(ws.cell(row,col['カート獲得アラート']).value or '')
        if cart_alert.startswith('異常'): paint(ws.cell(row,col['カート獲得アラート']),'red',True)
        elif cart_alert.startswith('要注意') or cart_alert.startswith('要確認'): paint(ws.cell(row,col['カート獲得アラート']),'orange')
        elif cart_alert.startswith('注意'): paint(ws.cell(row,col['カート獲得アラート']),'yellow')
        elif cart_alert=='問題なし': paint(ws.cell(row,col['カート獲得アラート']),'lightgreen')
        else: paint(ws.cell(row,col['カート獲得アラート']),'gray')

        margin=ws.cell(row,col['利益率']).value
        if isinstance(margin,(int,float)):
            if margin>=.10: paint(ws.cell(row,col['利益率']),'darkgreen',True)
            elif margin>=.05: paint(ws.cell(row,col['利益率']),'lightgreen')
            elif margin>=0: paint(ws.cell(row,col['利益率']),'yellow')
            else: paint(ws.cell(row,col['利益率']),'red',True)

    dash=wb.create_sheet('サマリー')
    dash.append(['指標','件数/数量'])
    dash.append(['DMD対象SKU',len(records)])
    dash.append(['最優先',sum(str(r['補充優先度']).startswith('1') for r in records)])
    dash.append(['至急',sum(str(r['補充優先度']).startswith('2') for r in records)])
    dash.append(['早め',sum(str(r['補充優先度']).startswith('3') for r in records)])
    dash.append(['発注判断',sum(r['仕入判断']=='発注' for r in records)])
    dash.append(['推奨発注数合計',sum(r['推奨発注数'] for r in records)])
    dash.append(['利益要確認',sum(r['仕入判断']=='利益要確認' for r in records)])
    dash.append(['30日以上滞留',sum(bool(r.get('滞留アラート')) for r in records)])
    dash.append(['滞留在庫金額',sum(r.get('現在庫金額',0) for r in records if r.get('滞留アラート'))])
    dash.append(['カート警告',sum(str(r.get('カート獲得アラート','')).startswith(('注意','要注意','異常','要確認')) for r in records)])
    dash.append(['最低利益率',f'{MIN_PROFIT_MARGIN:.1%}'])
    dash.append(['仕入戦略','売上拡大モード'])
    dash.append(['売上基準日',end.date().isoformat()])
    for c in dash[1]: c.fill=PatternFill('solid',fgColor=blue); c.font=Font(bold=True,color='FFFFFF')
    dash.column_dimensions['A'].width=28; dash.column_dimensions['B'].width=20

    raw=wb.create_sheet('分析詳細')
    detail_headers=['SKU','ASIN','商品名','現在庫','現在庫金額','30日販売','90日販売','最終販売日','無販売日数','滞留アラート','セッション数','カート獲得率','前60日月換算販売','販売速度比','カート獲得アラート','45日需要予測','SalesRank現在','Rank30日平均','Rank90日平均','新品価格','新品90日平均','FBA競合','FBM競合','Amazon本体','FBA最安値','新品最安値','採用価格種別','Keepaエラー']
    raw.append(detail_headers)
    for r in records:
        raw.append([r.get('SKU'),r.get('ASIN'),r.get('商品名'),r.get('現在庫'),r.get('現在庫金額'),r.get('30日販売'),r.get('90日販売'),r.get('最終販売日'),('90日以上' if r.get('最終販売日')=='90日以内なし' else r.get('無販売日数')),r.get('滞留アラート'),r.get('セッション数'),(r.get('カート獲得率')/100 if r.get('カート獲得率') is not None else None),r.get('前60日月換算販売'),r.get('販売速度比'),r.get('カート獲得アラート'),r.get('45日需要予測'),r.get('rank_now'),r.get('rank_avg30'),r.get('rank_avg90'),r.get('new_price'),r.get('new_avg90'),r.get('fba_count'),r.get('fbm_count'),r.get('amazon_offers') or r.get('amazon_present'),r.get('lowest_fba'),r.get('lowest_new'),r.get('価格種別'),r.get('keepa_error') or r.get('offers_error')])
    for c in raw[1]: c.fill=PatternFill('solid',fgColor=blue); c.font=Font(bold=True,color='FFFFFF')
    raw.freeze_panes='A2'; raw.auto_filter.ref=raw.dimensions

    if output_path is None:
        output_path=unique_output_path(RESULT_DIR, f'DMD仕入れ判断_{end.strftime("%Y%m%d")}.xlsx')
    else:
        os.makedirs(str(Path(output_path).parent), exist_ok=True)
    wb.save(output_path)
    print('Google Driveへ保存完了:',output_path)
    return output_path

if __name__=='__main__':
    main()
